"""Retain source events without inventing ownership or collapsing property history."""

import io
from datetime import date
from unittest.mock import AsyncMock

import openpyxl
import pytest

from app.crm import owner_queries
from app.crm.cadastral import _select_plot, _villa_code_variants, villa_sales_for_plots
from app.crm.identity import _candidate_owner_key, _claim_preference
from app.crm.import_events import event_family
from app.crm.import_models import CrmPropertyContactClaim
from app.crm.import_parsing import _inspect_csv, _inspect_xlsx, _parse_csv, _parse_xlsx
from app.crm.ownership import (
    _apply_evidence,
    _recover_property_identity_from_near_transaction,
    _recover_property_identity_from_transaction,
)


def parse(text):
    data = text.encode()
    dropped = []
    candidates, total = _parse_csv(data, "owners.csv", _inspect_csv(data, "owners.csv"), dropped)
    assert len(candidates) + len(dropped) == total
    return candidates


@pytest.mark.parametrize(
    "procedure,family",
    [
        ("Mortgage Registration", "financing"),
        ("Grant on Delayed Sell", "gift"),
        ("Complete Lease To Own On Delayed Sell", "lease_to_own"),
        ("Heirs Ownership", "inheritance"),
        ("Ownership Transfer", "ownership_transfer"),
        ("Property Annexation", "property_change"),
        ("Delayed Sell", "sale"),
        ("Novel procedure", "unknown"),
    ],
)
def test_procedures_retained_without_incidental_sale_inference(procedure, family):
    candidates = parse(
        "NameEn,ProcedurePartyTypeNameEn,ProcedureNameEn,BuildingNameEn,UnitNumber,Regis,ProcedureValue\n"
        f"Example,Seller,{procedure},Tower,101,2025-01-01,2000000\n"
    )
    row = candidates[0]
    assert event_family(procedure) == family
    _apply_evidence(
        row,
        sales={
            row.property_key: [
                {
                    "transaction_date": date(2025, 1, 1),
                    "sale_amount_aed": 2000000,
                }
            ]
        },
        rentals={},
    )
    assert row.ownership_status == (
        "former_owner" if family == "sale" else "unit_linked_unverified"
    )


def test_dated_owner_snapshot_and_multiple_phones_survive_without_buyer_role():
    rows = parse(
        "Owner Name,Date,BuildingName 2,UnitNumber,Transaction Amount,Phone 1,Mobile 1,Mobile 2,Secondary Mobile\n"
        "Example,2025-01-01,Tower,101,2000000,0,0501234567,0507654321,0501112233\n"
    )
    row = rows[0]
    assert row.transaction_price_aed == 2000000
    assert row.phone_primary and row.phone_alternate
    assert len(row.evidence["contact_channels"]) == 3
    _apply_evidence(
        row,
        sales={
            row.property_key: [
                {
                    "transaction_date": date(2025, 1, 1),
                    "sale_amount_aed": 2000000,
                }
            ]
        },
        rentals={},
    )
    # The matching DLD sale confirms the owner-list acquisition.
    assert row.ownership_status == "verified_current_owner"
    assert row.evidence["acquisition_inferred"] is True


def test_missing_unit_retains_land_evidence_without_zero_or_shared_owner_identity():
    rows = parse(
        "NameEn,Project,Building No,LandNumber,UnitNumber\n"
        "Example,Project,0,100,\nExample,Project,1,200,\n"
    )
    assert len(rows) == 2
    assert [r.evidence["source_identifiers"]["land_number"] for r in rows] == ["100", "200"]
    assert all(not r.unit_number for r in rows)
    assert _candidate_owner_key(rows[0]) != _candidate_owner_key(rows[1])


def test_formatted_tail_does_not_hide_data_after_a_blank_gap():
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["Owner Name", "Date", "Building No", "LandNumber"])
    sheet.append(["Example", "2025-01-01", "V-101", "12"])
    sheet.cell(100, 1, "Other")
    sheet.cell(1048576, 1).number_format = "0.00"
    out = io.BytesIO()
    book.save(out)
    profiles = _inspect_xlsx(out.getvalue(), "owners.xlsx")
    assert profiles[0].meaningful_rows == 2
    assert profiles[0].last_data_row == 100
    rows, total = _parse_xlsx(out.getvalue(), "owners.xlsx", profiles)
    assert total == len(rows) == 2
    assert rows[1].source_row_number == 100


def test_land_identity_requires_location_and_rejects_subdivision_collisions():
    candidate = parse("NameEn,Master Project,Project,LandNumber\nExample,Master,Cluster,12\n")[0]
    plot = {
        "land_number": "12",
        "land_sub_number": "0",
        "area_id": 5,
        "area_en": "Area",
        "master_project_en": "Master",
        "project_en": "Cluster",
    }
    assert _select_plot(candidate, [plot]) == plot
    assert _select_plot(candidate, [plot, {**plot, "area_id": 6}]) is None
    assert _select_plot(candidate, [plot, {**plot, "land_sub_number": "1"}]) is None
    assert _select_plot(candidate, [{**plot, "project_en": "Other"}]) is None
    candidate.evidence["source_fields"] = {}
    assert _select_plot(candidate, [plot]) is None


def test_later_reacquisition_beats_old_disposal_but_new_mortgage_does_not():
    old_sale = CrmPropertyContactClaim(
        ownership_status="former_owner", confidence_score=100, transaction_date=date(2020, 1, 1)
    )
    purchase = CrmPropertyContactClaim(
        ownership_status="verified_current_owner",
        confidence_score=98,
        transaction_date=date(2024, 1, 1),
    )
    mortgage = CrmPropertyContactClaim(
        ownership_status="unit_linked_unverified",
        confidence_score=50,
        transaction_date=date(2025, 1, 1),
    )
    assert max([old_sale, purchase, mortgage], key=_claim_preference) is purchase


async def test_financing_amount_is_never_looked_up_as_sale_signature(monkeypatch):
    rows = parse(
        "NameEn,Regis,ProcedureValue,ProcedureNameEn,ProcedurePartyTypeNameEn\n"
        "Example,2025-01-01,2000000,Mortgage Registration,Buyer\n"
    )
    # Both signature resolvers must exit before any warehouse request.
    query = AsyncMock(side_effect=AssertionError("A financing amount reached sales matching"))
    monkeypatch.setattr(owner_queries, "query", query)
    await _recover_property_identity_from_transaction(rows)
    await _recover_property_identity_from_near_transaction(rows)
    query.assert_not_awaited()


def test_villa_code_variants_cover_prefixed_and_bare_unit_numbers():
    # DLD sales record the same villa as "DH2-XV728B" and as "XV728B".
    assert _villa_code_variants("DH2-XV728B") == {"dh2-xv728b", "xv728b"}
    assert _villa_code_variants(" xv 728b ") == {"xv728b"}
    assert _villa_code_variants("0") == set()


async def test_villa_sales_link_plot_to_its_own_community_only(monkeypatch):
    # A plot's pre-registration code identifies its villa; the same code in another
    # community (e.g. Damac Hills 1) is a different property and must be ignored.
    row = parse("NameEn,Master Project,Project,LandNumber\nExample,DAMAC HILLS 2,0,2822\n")[0]
    row.evidence["cadastral_identity"] = {
        "property_key": "dld-land:505:2822:0",
        "villa_code": "DH2-XV728B",
        "master_project": "DAMAC HILLS 2",
        "project": "DAMAC HILLS (2) - BASSWOOD",
        "area_name": "Madinat Hind 4",
    }

    def sale(unit_key, area, on):
        return {
            "unit_key": unit_key,
            "area_name": area,
            "project_name": "",
            "transaction_date": on,
            "sale_amount_aed": 1_200_000,
            "sale_key": f"{unit_key}-{on}",
        }

    fetch = AsyncMock(
        return_value=[
            sale("dh2-xv728b", "Damac Hills 2", date(2016, 11, 29)),
            sale("xv728b", "Damac Hills 2", date(2024, 2, 13)),
            sale("xv728b", "Damac Hills", date(2018, 6, 27)),
        ]
    )
    monkeypatch.setattr(owner_queries, "fetch_unit_sales", fetch)

    sales = await villa_sales_for_plots([row])

    assert set(fetch.await_args.args[0]["unit_keys"]) == {"dh2-xv728b", "xv728b"}
    assert [s["transaction_date"] for s in sales["dld-land:505:2822:0"]] == [
        date(2016, 11, 29),
        date(2024, 2, 13),
    ]
