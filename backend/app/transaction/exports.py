import csv
from datetime import UTC, datetime
from io import BytesIO, StringIO
from typing import Literal, Optional

from openpyxl import Workbook
from openpyxl.styles import Font

from app.transaction import queries

MAX_TRANSACTION_EXPORT_ROWS = 50_000

SPREADSHEET_FORMULA_PREFIXES = ("=", "+", "-", "@")

SPREADSHEET_CONTROL_PREFIXES = ("\t", "\r", "\n")


def _active_export_filters(filters: dict[str, object]) -> str:
    active_filters = [
        f"{key}={value}"
        for key, value in filters.items()
        if value is not None and value != "" and value != "All"
    ]
    return ", ".join(active_filters) if active_filters else "None"


def _add_export_metadata(
    workbook: Workbook,
    *,
    title: str,
    generated_at: datetime,
    row_count: int,
    filters: dict[str, object],
    source: str,
) -> None:
    metadata = workbook.create_sheet("Metadata", 0)
    rows = [
        ("Publisher", "Dubai Land Department"),
        ("Dataset", title),
        ("Generated at", generated_at.isoformat()),
        ("Data source", source),
        ("Rows exported", row_count),
        ("Maximum filtered rows", MAX_TRANSACTION_EXPORT_ROWS),
        ("Filters applied", _active_export_filters(filters)),
    ]
    for row in rows:
        metadata.append(row)
    for cell in metadata["A"]:
        cell.font = Font(bold=True)


def _add_export_header(
    worksheet,
    *,
    title: str,
    generated_at: datetime,
    source: str,
) -> None:
    worksheet.append(["Dubai Land Department"])
    worksheet.append([title])
    worksheet.append([f"Generated at {generated_at.isoformat()}"])
    worksheet.append([f"Source: {source}"])
    worksheet.append([])
    worksheet["A1"].font = Font(bold=True, size=14)
    worksheet["A2"].font = Font(bold=True)


def _sanitize_spreadsheet_value(value: object) -> object:
    if not isinstance(value, str) or not value:
        return value

    stripped = value.lstrip()
    if value.startswith(SPREADSHEET_CONTROL_PREFIXES) or (
        stripped and stripped.startswith(SPREADSHEET_FORMULA_PREFIXES)
    ):
        return f"'{value}"
    return value


def _sanitize_spreadsheet_row(row: list[object]) -> list[object]:
    return [_sanitize_spreadsheet_value(value) for value in row]


def _rows_to_csv(headers: list[str], rows: list[list[object]]) -> BytesIO:
    csv_text = StringIO()
    writer = csv.writer(csv_text)
    writer.writerow(headers)
    writer.writerows(_sanitize_spreadsheet_row(row) for row in rows)

    output = BytesIO()
    output.write(csv_text.getvalue().encode("utf-8-sig"))
    output.seek(0)
    return output


async def export_sales_transactions(
    transaction_ids: Optional[str] = None,
    property_usage: Optional[str] = None,
    property_type: Optional[str] = None,
    property_sub_type: Optional[str] = None,
    rooms: Optional[str] = None,
    registration_type: Optional[str] = None,
    area: Optional[str] = None,
    project: Optional[str] = None,
    filter_type: str = "project",
    building: Optional[str] = None,
    timeframe_days: Optional[int] = None,
    area_min: Optional[float] = None,
    area_max: Optional[float] = None,
    price_min: Optional[float] = None,
    price_max: Optional[float] = None,
    price_per_sqm_min: Optional[float] = None,
    price_per_sqm_max: Optional[float] = None,
    file_format: Literal["xlsx", "csv"] = "xlsx",
) -> BytesIO:
    """
    Export sales transactions to Excel or CSV.

    Limit: 50,000 rows max (or specific transaction_ids if provided).
    """
    filters = {
        "transaction_ids": transaction_ids,
        "property_usage": property_usage,
        "property_type": property_type,
        "property_sub_type": property_sub_type,
        "rooms": rooms,
        "registration_type": registration_type,
        "area": area,
        "project": project,
        "filter_type": filter_type,
        "building": building,
        "timeframe_days": timeframe_days,
        "area_min": area_min,
        "area_max": area_max,
        "price_min": price_min,
        "price_max": price_max,
        "price_per_sqm_min": price_per_sqm_min,
        "price_per_sqm_max": price_per_sqm_max,
    }

    # If specific IDs provided, use them (no limit)
    results = await queries.export_sales_transactions_rows(
        MAX_TRANSACTION_EXPORT_ROWS,
        area,
        area_max,
        area_min,
        building,
        filter_type,
        price_max,
        price_min,
        price_per_sqm_max,
        price_per_sqm_min,
        project,
        property_sub_type,
        property_type,
        property_usage,
        registration_type,
        rooms,
        timeframe_days,
        transaction_ids,
    )

    headers = [
        "Transaction ID",
        "Date",
        "Usage",
        "Trans Group",
        "Procedure",
        "Type",
        "Sub Type",
        "Rooms",
        "Area (sqm)",
        "Price",
        "Price/sqm",
        "Registration Type",
        "Area",
        "Project",
        "Building",
        "Landmark",
        "Metro",
        "Mall",
    ]
    rows = [
        [
            row["transaction_id"],
            row["instance_date"],
            row["property_usage_en"],
            row["trans_group_en"],
            row["procedure_name_en"],
            row["property_type_en"],
            row["property_sub_type_en"],
            row["rooms_en"],
            float(row["procedure_area"]) if row["procedure_area"] else None,
            float(row["actual_worth"]) if row["actual_worth"] else None,
            float(row["meter_sale_price"]) if row["meter_sale_price"] else None,
            row["reg_type_en"],
            row["area_name_en"],
            row["project_name_en"],
            row["building_name_en"],
            row["nearest_landmark_en"],
            row["nearest_metro_en"],
            row["nearest_mall_en"],
        ]
        for row in results
    ]

    if file_format == "csv":
        return _rows_to_csv(headers, rows)

    generated_at = datetime.now(UTC)
    source = "DLD transaction registry"

    # Create Excel workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales Transactions"
    _add_export_metadata(
        wb,
        title="Sales Transactions",
        generated_at=generated_at,
        row_count=len(rows),
        filters=filters,
        source=source,
    )
    _add_export_header(ws, title="Sales Transactions", generated_at=generated_at, source=source)
    ws.append(headers)

    # Style headers
    for cell in ws[6]:
        cell.font = Font(bold=True)

    # Add data rows
    for row in rows:
        ws.append(_sanitize_spreadsheet_row(row))

    # Save to BytesIO
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


async def export_rental_contracts(
    contract_ids: Optional[str] = None,
    property_usage: Optional[str] = None,
    property_type: Optional[str] = None,
    property_sub_type: Optional[str] = None,
    tenant_type: Optional[str] = None,
    contract_reg_type: Optional[str] = None,
    area: Optional[str] = None,
    project: Optional[str] = None,
    filter_type: str = "project",
    timeframe_days: Optional[int] = None,
    area_min: Optional[float] = None,
    area_max: Optional[float] = None,
    rent_min: Optional[float] = None,
    rent_max: Optional[float] = None,
    file_format: Literal["xlsx", "csv"] = "xlsx",
) -> BytesIO:
    """
    Export rental contracts to Excel or CSV.

    Limit: 50,000 rows max (or specific contract_ids if provided).
    """
    filters = {
        "contract_ids": contract_ids,
        "property_usage": property_usage,
        "property_type": property_type,
        "property_sub_type": property_sub_type,
        "tenant_type": tenant_type,
        "contract_reg_type": contract_reg_type,
        "area": area,
        "project": project,
        "filter_type": filter_type,
        "timeframe_days": timeframe_days,
        "area_min": area_min,
        "area_max": area_max,
        "rent_min": rent_min,
        "rent_max": rent_max,
    }

    # If specific IDs provided, use them (no limit)
    results = await queries.export_rental_contracts_rows(
        MAX_TRANSACTION_EXPORT_ROWS,
        area,
        area_max,
        area_min,
        contract_ids,
        contract_reg_type,
        filter_type,
        project,
        property_sub_type,
        property_type,
        property_usage,
        rent_max,
        rent_min,
        tenant_type,
        timeframe_days,
    )

    headers = [
        "Contract ID",
        "Start Date",
        "End Date",
        "Usage",
        "Reg Type",
        "Type",
        "Sub Type",
        "Tenant Type",
        "Area (sqm)",
        "Annual Rent",
        "Contract Amount",
        "Area",
        "Project",
        "Landmark",
        "Metro",
        "Mall",
        "No. Properties",
        "Freehold",
    ]
    rows = [
        [
            row["contract_id"],
            row["contract_start_date"],
            row["contract_end_date"],
            row["property_usage_en"],
            row["contract_reg_type_en"],
            row["ejari_property_type_en"],
            row["ejari_property_sub_type_en"],
            row["tenant_type_en"],
            float(row["actual_area"]) if row["actual_area"] else None,
            float(row["annual_amount"]) if row["annual_amount"] else None,
            float(row["contract_amount"]) if row["contract_amount"] else None,
            row["area_name_en"],
            row["project_name_en"],
            row["nearest_landmark_en"],
            row["nearest_metro_en"],
            row["nearest_mall_en"],
            row["no_of_prop"],
            "Yes" if row["is_free_hold"] else "No",
        ]
        for row in results
    ]

    if file_format == "csv":
        return _rows_to_csv(headers, rows)

    generated_at = datetime.now(UTC)
    source = "DLD Ejari rental contracts"

    # Create Excel workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Rental Contracts"
    _add_export_metadata(
        wb,
        title="Rental Contracts",
        generated_at=generated_at,
        row_count=len(rows),
        filters=filters,
        source=source,
    )
    _add_export_header(ws, title="Rental Contracts", generated_at=generated_at, source=source)
    ws.append(headers)

    # Style headers
    for cell in ws[6]:
        cell.font = Font(bold=True)

    # Add data rows
    for row in rows:
        ws.append(_sanitize_spreadsheet_row(row))

    # Save to BytesIO
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output
