from __future__ import annotations

import csv
import hashlib
import io
import json
from datetime import timedelta
from pathlib import Path
from typing import Any

import openpyxl
from fastapi import HTTPException, UploadFile
from sqlmodel.ext.asyncio.session import AsyncSession

from app.crm import inbound_repository
from app.crm.ai import _ai_extract, _ai_header_mapping
from app.crm.identity import _normalize_email, _normalize_phone
from app.crm.inbound import _bulk_ownership_enrichments, _find_or_create_contact
from app.crm.inbound_models import (
    CrmInboundActivity,
    CrmInboundContact,
    CrmInboundEnquiry,
    CrmInboundImport,
    InboundImportPreviewResponse,
    InboundImportPreviewSheet,
    InboundImportResponse,
    InboundImportSummary,
    InboundOwnershipEnrichmentResponse,
)
from app.crm.inbound_policy import (
    _ALIASES,
    _FIRST_RESPONSE_MINUTES,
    _deterministic_criteria,
    _lead_priority,
    _normalize_intent,
    _parse_datetime,
    _redact,
    _text,
    _token,
    _utcnow,
)

_MAX_FILE_BYTES = 50 * 1024 * 1024


def _header_mapping(headers: list[Any]) -> dict[str, int]:
    tokens = {_token(header): index for index, header in enumerate(headers) if _text(header)}
    mapping: dict[str, int] = {}
    for canonical, aliases in _ALIASES.items():
        for alias in aliases:
            if alias in tokens:
                mapping[canonical] = tokens[alias]
                break
    return mapping


def _best_header(rows: list[tuple[Any, ...]]) -> tuple[int, tuple[Any, ...], dict[str, int]]:
    """Find the most likely header row in exports with titles or blank preambles."""
    candidates = []
    for index, values in enumerate(rows[:15]):
        mapping = _header_mapping(list(values))
        identity_fields = len({"full_name", "first_name", "phone", "email"} & mapping.keys())
        candidates.append((len(mapping), identity_fields, -index, index, values, mapping))
    if not candidates:
        return 0, (), {}
    _, _, _, index, values, mapping = max(candidates, key=lambda item: item[:3])
    return index, values, mapping


def _mapping_override_for_sheet(
    override: dict[str, Any] | None,
    sheet_name: str,
    headers: list[Any],
    detected: dict[str, int],
) -> dict[str, int]:
    if not override:
        return detected
    sheets = override.get("sheets", [])
    sheet_config = next((item for item in sheets if item.get("sheet") == sheet_name), None)
    if not sheet_config:
        return detected
    by_header = {_text(header): index for index, header in enumerate(headers)}
    mapped = {
        canonical: by_header[header]
        for canonical, header in sheet_config.get("field_map", {}).items()
        if header in by_header and canonical in _ALIASES
    }
    return mapped or detected


def _file_rows(
    data: bytes,
    filename: str,
    mapping_override: dict[str, Any] | None = None,
) -> tuple[list[tuple[str, int, dict[str, Any]]], dict[str, Any]]:
    rows: list[tuple[str, int, dict[str, Any]]] = []
    mappings: list[dict[str, Any]] = []
    if filename.lower().endswith(".csv"):
        raw = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
        header_index, header_values, mapping = _best_header([tuple(row) for row in raw])
        headers = list(header_values)
        sheet_name = Path(filename).stem
        mapping = _mapping_override_for_sheet(mapping_override, sheet_name, headers, mapping)
        for row_number, values in enumerate(raw[header_index + 1 :], start=header_index + 2):
            if any(_text(value) for value in values):
                rows.append(
                    (
                        sheet_name,
                        row_number,
                        {
                            key: values[index] if index < len(values) else None
                            for key, index in mapping.items()
                        },
                    )
                )
        mappings.append(
            {
                "sheet": sheet_name,
                "header_row": header_index + 1,
                "headers": [_text(value) for value in headers if _text(value)],
                "field_map": {key: headers[index] for key, index in mapping.items()},
            }
        )
        return rows, {"sheets": mappings}

    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            preview = list(sheet.iter_rows(min_row=1, max_row=15, values_only=True))
            header_index, header_values, mapping = _best_header(preview)
            mapping = _mapping_override_for_sheet(
                mapping_override, sheet.title, list(header_values), mapping
            )
            mappings.append(
                {
                    "sheet": sheet.title,
                    "header_row": header_index + 1,
                    "headers": [_text(value) for value in header_values if _text(value)],
                    "field_map": {
                        key: _text(header_values[index]) for key, index in mapping.items()
                    },
                }
            )
            for row_number, values in enumerate(
                sheet.iter_rows(min_row=header_index + 2, values_only=True), start=header_index + 2
            ):
                if any(value not in (None, "") for value in values):
                    rows.append(
                        (
                            sheet.title,
                            row_number,
                            {
                                key: values[index] if index < len(values) else None
                                for key, index in mapping.items()
                            },
                        )
                    )
    finally:
        workbook.close()
    return rows, {"sheets": mappings}


def _preview_file(
    data: bytes,
    filename: str,
    mapping_override: dict[str, Any] | None = None,
    ai_mapped_sheets: set[str] | None = None,
) -> InboundImportPreviewResponse:
    rows, audit = _file_rows(data, filename, mapping_override)
    grouped: dict[str, list[tuple[int, dict[str, Any]]]] = {}
    for sheet, row_number, row in rows:
        grouped.setdefault(sheet, []).append((row_number, row))
    sheets = []
    for item in audit.get("sheets", []):
        sheet_name = item["sheet"]
        field_map = item.get("field_map", {})
        headers = item.get("headers", [])
        samples = []
        for row_number, row in grouped.get(sheet_name, [])[:5]:
            samples.append(
                {
                    "row": str(row_number),
                    **{
                        key: _redact(_text(value), [])[:160]
                        for key, value in row.items()
                        if value not in (None, "") and key not in {"phone", "email", "whatsapp"}
                    },
                    **{key: "Present" for key in ("phone", "email", "whatsapp") if row.get(key)},
                }
            )
        sheets.append(
            InboundImportPreviewSheet(
                sheet=sheet_name,
                header_row=item.get("header_row", 1),
                headers=headers,
                field_map=field_map,
                sample_rows=samples,
                row_count=len(grouped.get(sheet_name, [])),
                mapping_source=(
                    "ai_assisted" if sheet_name in (ai_mapped_sheets or set()) else "deterministic"
                ),
            )
        )
    warnings = []
    if not any("full_name" in sheet.field_map for sheet in sheets):
        warnings.append("No name column was detected")
    if not any({"phone", "email"} & set(sheet.field_map) for sheet in sheets):
        warnings.append("No phone or email column was detected")
    return InboundImportPreviewResponse(
        original_filename=filename,
        size_bytes=len(data),
        total_rows=len(rows),
        sheets=sheets,
        warnings=warnings,
    )


async def preview_inbound_import(
    upload: UploadFile, campaign_context: str = ""
) -> InboundImportPreviewResponse:
    filename = Path(upload.filename or "inbound-leads.xlsx").name
    if not filename.lower().endswith((".xlsx", ".csv")):
        raise HTTPException(status_code=415, detail="Upload a CSV or XLSX file")
    data = await upload.read(_MAX_FILE_BYTES + 1)
    if len(data) > _MAX_FILE_BYTES:
        raise HTTPException(status_code=413, detail="Inbound import exceeds 50 MB")
    try:
        preview = _preview_file(data, filename)
        ai_mapped_sheets: set[str] = set()
        for sheet in preview.sheets:
            has_identity = bool(
                {"full_name", "first_name", "phone", "email"} & set(sheet.field_map)
            )
            if has_identity and len(sheet.field_map) >= 4:
                continue
            suggested = await _ai_header_mapping(sheet.headers, campaign_context.strip()[:4_000])
            for canonical, header in suggested.items():
                sheet.field_map.setdefault(canonical, header)
            if suggested:
                ai_mapped_sheets.add(sheet.sheet)
        if ai_mapped_sheets:
            override = {
                "sheets": [
                    {"sheet": sheet.sheet, "field_map": sheet.field_map} for sheet in preview.sheets
                ]
            }
            preview = _preview_file(data, filename, override, ai_mapped_sheets)
            preview.warnings.append(
                "AI suggested mappings are highlighted for review and are not applied until approval"
            )
        return preview
    except Exception as exc:
        raise HTTPException(status_code=422, detail=f"Could not read import: {exc}") from exc


async def process_inbound_import(
    db: AsyncSession,
    *,
    upload: UploadFile,
    name: str,
    user_id: str,
    organization_id: str,
    campaign_context: str = "",
    mapping_override: dict[str, Any] | None = None,
) -> InboundImportResponse:
    filename = Path(upload.filename or "inbound-leads.xlsx").name
    if not filename.lower().endswith((".xlsx", ".csv")):
        raise HTTPException(status_code=415, detail="Upload a CSV or XLSX file")
    data = await upload.read(_MAX_FILE_BYTES + 1)
    if len(data) > _MAX_FILE_BYTES:
        raise HTTPException(status_code=413, detail="Inbound import exceeds 50 MB")
    campaign_context = campaign_context.strip()[:4_000]
    digest = hashlib.sha256(data + b"\0" + campaign_context.casefold().encode()).hexdigest()
    existing_result = await inbound_repository.find_import_by_digest(db, digest, organization_id)
    if existing := existing_result:
        return _import_response(existing)

    now = _utcnow()
    record = CrmInboundImport(
        organization_id=organization_id,
        created_by_user_id=user_id,
        name=name.strip() or Path(filename).stem,
        original_filename=filename,
        sha256=digest,
        size_bytes=len(data),
        status="processing",
        mapping={},
        summary={},
        created_at=now,
    )
    inbound_repository.save(db, record)
    await db.commit()
    await db.refresh(record)
    summary = InboundImportSummary()
    imported_enquiries: list[CrmInboundEnquiry] = []
    imported_contacts: dict[int, CrmInboundContact] = {}
    try:
        rows, mapping = _file_rows(data, filename, mapping_override)
        mapping["campaign_context"] = campaign_context
        record.mapping = mapping
        summary.total_rows = len(rows)
        for sheet_mapping in mapping.get("sheets", []):
            fields = set(sheet_mapping.get("field_map", {}))
            if not ({"full_name", "first_name"} & fields):
                raise ValueError(
                    f"Map a full-name or first-name column for {sheet_mapping.get('sheet', 'the import')}"
                )
            if not ({"phone", "email", "whatsapp"} & fields):
                raise ValueError(
                    f"Map a phone, email, or WhatsApp column for {sheet_mapping.get('sheet', 'the import')}"
                )
        if summary.total_rows > 10_000:
            raise ValueError("Inbound import exceeds the 10,000-row synchronous import limit")
        row_fingerprints: set[str] = set()
        for sheet, row_number, row in rows:
            first = _text(row.get("first_name"))
            last = _text(row.get("last_name"))
            full_name = _text(row.get("full_name")) or " ".join(
                part for part in (first, last) if part
            )
            phone = _normalize_phone(row.get("phone"))
            email = _normalize_email(row.get("email"))
            if not any((full_name, phone, email)):
                continue
            source_id = _text(row.get("source_lead_id"))
            fingerprint_payload = {
                "source_id": source_id,
                "name": full_name.casefold(),
                "phone": phone,
                "email": email,
                "intent": _text(row.get("intent")).casefold(),
                "notes": _text(row.get("notes")).casefold(),
                "locations": _text(row.get("locations")).casefold(),
                "budget": _text(row.get("budget_max")).casefold(),
                "created_at": _text(row.get("created_at")),
            }
            row_fingerprint = hashlib.sha256(
                json.dumps(fingerprint_payload, sort_keys=True).encode()
            ).hexdigest()
            if row_fingerprint in row_fingerprints:
                summary.duplicates += 1
                continue
            row_fingerprints.add(row_fingerprint)
            notes = _text(row.get("notes"))
            # Campaign context describes provenance, not an individual's intent.
            # Only row fields and the person's message may establish observed facts.
            row_context = notes
            intent = _normalize_intent(row.get("intent"), row_context)
            criteria = _deterministic_criteria(row, row_context)
            redacted = _redact(row_context, [full_name, phone or "", email or ""])
            needs_ai = bool(row_context) and (
                intent is None
                or not criteria.get("locations")
                or (
                    intent in {"buy", "rent"}
                    and not any(
                        criteria.get(field) is not None
                        for field in ("bedrooms_min", "budget_max_aed", "property_type")
                    )
                )
            )
            ai_values, confidence, strict = (
                await _ai_extract(redacted, criteria) if needs_ai else ({}, 1.0, [])
            )
            ai_intent = ai_values.pop("_intent", None)
            review: list[str] = []
            contact = await _find_or_create_contact(
                db,
                organization_id=organization_id,
                full_name=full_name,
                phone=phone,
                email=email,
                whatsapp=_normalize_phone(row.get("whatsapp")),
                language=_text(row.get("language")) or None,
                consent_status=_text(row.get("consent")) or "unknown",
            )
            inferred_data: dict[str, Any] = {}
            if ai_intent in {"buy", "rent", "sell", "let"} and intent is None:
                inferred_data["intent"] = ai_intent
            if ai_values:
                inferred_data["criteria"] = ai_values
            if intent is None:
                review.append("Intent was not explicitly supplied and must be qualified")
            if intent in {"buy", "rent"} and not criteria.get("locations"):
                review.append("Preferred location is missing")
            if confidence and confidence < 0.85:
                review.append("Message extraction should be reviewed")
            record_kind = "opportunity" if intent else "unqualified_contact"
            observed_data = {
                "intent": intent,
                "message": notes or None,
                "criteria": criteria,
                "source": _text(row.get("source")) or "Imported",
                "campaign": _text(row.get("campaign")) or record.name,
                "campaign_context": campaign_context or None,
            }
            enquiry_date = _parse_datetime(row.get("created_at"))
            enquiry = CrmInboundEnquiry(
                organization_id=organization_id,
                contact_id=contact.id or 0,
                source_import_id=record.id,
                created_by_user_id=user_id,
                assigned_user_id=user_id,
                record_kind=record_kind,
                intent=intent or "unknown",
                status="needs_review" if review else "new",
                source=_text(row.get("source")) or "Imported",
                campaign=_text(row.get("campaign")) or record.name,
                source_lead_id=source_id or None,
                source_sheet=sheet,
                source_row_number=row_number,
                enquiry_date=enquiry_date,
                criteria=criteria,
                strict_fields=strict,
                raw_notes=notes or None,
                observed_data=observed_data,
                inferred_data=inferred_data,
                confirmed_data={},
                qualification={},
                extraction_confidence=confidence or (0.8 if review else 0.9),
                extraction_evidence={
                    "redacted_ai": bool(ai_values),
                    "campaign_context_used": bool(campaign_context),
                    "classification_source": "row_or_message" if intent else "unknown",
                    "mapping": mapping,
                },
                review_reasons=review,
                first_response_due_at=enquiry_date + timedelta(minutes=_FIRST_RESPONSE_MINUTES),
                created_at=now,
                updated_at=now,
            )
            inbound_repository.save(db, enquiry)
            await db.flush()
            imported_enquiries.append(enquiry)
            imported_contacts[contact.id or 0] = contact
            inbound_repository.save(
                db,
                CrmInboundActivity(
                    enquiry_id=enquiry.id or 0,
                    organization_id=organization_id,
                    user_id=user_id,
                    activity_type="imported",
                    body=f"Imported from {record.name}",
                    metadata_json={"sheet": sheet, "row": row_number},
                    created_at=now,
                ),
            )
            summary.imported += 1
            summary.needs_review += int(bool(review))
            if intent == "buy":
                summary.buyers += 1
            elif intent == "rent":
                summary.tenants += 1
            elif intent == "sell":
                summary.sellers += 1
            elif intent == "let":
                summary.landlords += 1
            else:
                summary.unqualified_contacts += 1
        enrichments = await _bulk_ownership_enrichments(
            db,
            organization_id=organization_id,
            contacts=list(imported_contacts.values()),
        )
        for enquiry in imported_enquiries:
            enrichment = enrichments.get(enquiry.contact_id, InboundOwnershipEnrichmentResponse())
            summary.owner_matches += int(enrichment.matched)
            summary.current_owners += int(enrichment.current_owner)
            summary.former_owners += int(enrichment.former_owner)
            summary.rental_owners += int(enrichment.rental_owner)
            summary.portfolio_owners += int(enrichment.multiple_property_owner)
            summary.high_priority += int(_lead_priority(enquiry, enrichment).band == "high")
        record.status = "completed"
        record.summary = summary.model_dump()
        record.completed_at = _utcnow()
        inbound_repository.save(db, record)
        await db.commit()
        await db.refresh(record)
    except Exception as exc:
        await db.rollback()
        record.status = "failed"
        record.error_message = str(exc)[:1_000]
        record.summary = summary.model_dump()
        record.completed_at = _utcnow()
        inbound_repository.save(db, record)
        await db.commit()
    return _import_response(record)


def _import_response(record: CrmInboundImport) -> InboundImportResponse:
    return InboundImportResponse(
        id=record.id or 0,
        name=record.name,
        original_filename=record.original_filename,
        status=record.status,  # type: ignore[arg-type]
        summary=InboundImportSummary.model_validate(record.summary or {}),
        mapping=record.mapping or {},
        error_message=record.error_message,
        created_at=record.created_at,
        completed_at=record.completed_at,
    )


async def list_inbound_imports(
    db: AsyncSession, organization_id: str
) -> list[InboundImportResponse]:
    result = await inbound_repository.list_imports(db, organization_id)
    return [_import_response(item) for item in result]


async def delete_inbound_import(db: AsyncSession, import_id: int, organization_id: str) -> None:
    result = await inbound_repository.get_import_for_org(db, import_id, organization_id)
    record = result
    if not record:
        raise HTTPException(status_code=404, detail="Inbound import not found")
    imported_enquiries = await inbound_repository.list_import_contact_ids(
        db, import_id, organization_id
    )
    contact_ids = list(dict.fromkeys(imported_enquiries))
    await inbound_repository.delete_import_enquiries(db, import_id, organization_id)
    if contact_ids:
        await inbound_repository.delete_orphan_contacts(db, contact_ids, organization_id)
    await inbound_repository.delete_record(db, record)
    await db.commit()
