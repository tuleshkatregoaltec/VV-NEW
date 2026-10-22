from __future__ import annotations

import csv
import hashlib
import io
import re
import xml.etree.ElementTree as ET
from collections.abc import Iterable
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl

from app.crm.identity import (
    _NON_KEY,
    ParsedContact,
    _normalize_email,
    _normalize_phone,
    _normalize_unit,
    _recover_embedded_property_identity,
    _text,
)
from app.crm.import_events import IMPORT_RESOLVER_VERSION, event_family
from app.crm.import_inventory import SourcePropertyInventory
from app.crm.import_models import (
    CrmContactImportSummary,
    CrmSheetMapping,
)

_MAX_ROWS = 500_000

_PLACEHOLDER_HEADER = re.compile(r"^column\s*\d+$", re.IGNORECASE)
_BUILDING_NO_PROPERTY_IDENTIFIER = re.compile(
    r"(?:^|[-\s])(?:villa|v|th)[-\s]*(?=[a-z0-9.]*\d)[a-z0-9.]+$", re.IGNORECASE
)
_BUILDING_NO_NUMERIC_IDENTIFIER = re.compile(
    r"^(?:r\d*\s+)?(.+?)[-\s]+(?:v|villa|th)[-\s]*(\d+)$", re.IGNORECASE
)

_CANONICAL_FIELDS = {
    "area_name",
    "master_project",
    "project_name",
    "building_name",
    "unit_number",
    "property_type",
    "completion_status",
    "size_sqft",
    "transaction_date",
    "transaction_price_aed",
    "party_role",
    "procedure_type",
    "contact_name",
    "phone_primary",
    "phone_alternate",
    "email_primary",
    "building_number",
    "land_number",
    "land_sub_number",
    "property_number",
    "plot_number",
    "municipality_number",
    "municipality_sub_number",
}

_ALIASES = {
    "area_name": {"areanameen", "area", "community", "masterlocation"},
    "master_project": {"masterproject", "masterdevelopment"},
    "project_name": {"project", "projectname", "development"},
    "building_name": {"buildingnameen", "building", "tower", "buildingname", "buildingname2"},
    "unit_number": {
        "unitnumber",
        "unit",
        "villanumber",
        "unitnumberproxy",
        "plotpreregisterationno",
        "plotpreregistrationno",
        "plotregistrationno",
    },
    "property_type": {"propertytypeen", "propertytype", "type"},
    "completion_status": {
        "completionstatus",
        "marketstatus",
        "registrationtype",
        "regtypeen",
        "propertycompletionstatus",
    },
    "size_sqft": {"size", "areasqft", "propertysize"},
    "transaction_date": {"regis", "date", "registrationdate", "transactiondate"},
    "transaction_price_aed": {
        "procedurevalue",
        "price",
        "transactionvalue",
        "saleprice",
        "transactionamount",
    },
    "party_role": {"procedurepartytypenameen", "partytype", "partyrole"},
    "procedure_type": {"procedurenameen", "procedurename"},
    "contact_name": {"nameen", "name", "ownername", "contactname"},
    "phone_primary": {"phone", "mobile", "phonenumber", "mobilenumber", "mobile1", "phone1"},
    "phone_alternate": {"secondarymobile", "alternatemobile", "mobile2"},
    "email_primary": {"email", "emailaddress", "owneremail"},
    "building_number": {"buildingno", "buildingnumber", "building1"},
    "land_number": {"landnumber", "landno"},
    "land_sub_number": {"landsubnumber", "landsubno"},
    "property_number": {"propertynumber"},
    "plot_number": {"plotpreregno", "plotpreregistrationno", "plotnumber"},
    "municipality_number": {"dmno", "municipalityno"},
    "municipality_sub_number": {"dmsubno", "municipalitysubno"},
}


@dataclass
class HeaderCell:
    key: str
    label: str
    index: int


@dataclass
class SheetProfile:
    name: str
    header_row: int
    headers: list[HeaderCell]
    meaningful_rows: int
    field_map: dict[str, str]
    sheet_family: str
    mapping_source: str = "heuristic"
    confidence: float = 0.82
    warnings: list[str] = field(default_factory=list)
    date_order: str = "dmy"
    last_data_row: int | None = None
    property_inventory: SourcePropertyInventory = field(default_factory=SourcePropertyInventory)


def _header_token(value: str) -> str:
    base = value.rsplit("__", 1)[0]
    return _NON_KEY.sub("", base.lower())


def _unique_headers(values: tuple[Any, ...] | list[Any]) -> list[HeaderCell]:
    counts: dict[str, int] = {}
    headers: list[HeaderCell] = []
    for index, raw in enumerate(values):
        label = _text(raw)
        if not label or _PLACEHOLDER_HEADER.match(label):
            continue
        count = counts.get(label.lower(), 0) + 1
        counts[label.lower()] = count
        key = label if count == 1 else f"{label}__{count}"
        headers.append(HeaderCell(key=key, label=label, index=index))
    return headers


def _recognized_headers(values: tuple[Any, ...] | list[Any]) -> int:
    tokens = {_header_token(_text(value)) for value in values if _text(value)}
    aliases = set().union(*_ALIASES.values())
    return len(tokens & aliases)


def _find_header(rows: list[tuple[Any, ...]]) -> tuple[int, list[HeaderCell]]:
    best_index = 0
    best_score = -1
    best_headers: list[HeaderCell] = []
    for index, row in enumerate(rows):
        headers = _unique_headers(row)
        score = _recognized_headers(row) * 100 + len(headers)
        if score > best_score:
            best_index = index
            best_score = score
            best_headers = headers
    return best_index + 1, best_headers


def _header_for(
    headers: list[HeaderCell],
    aliases: set[str],
    *,
    last: bool = False,
) -> str | None:
    matches = [header.key for header in headers if _header_token(header.key) in aliases]
    if not matches:
        return None
    return matches[-1] if last else matches[0]


def _heuristic_mapping(
    headers: list[HeaderCell],
    *,
    filename: str,
    sheet_name: str,
) -> tuple[dict[str, str], str, list[str]]:
    mapping: dict[str, str] = {}
    warnings: list[str] = []
    lowered_filename = filename.lower()
    for field_name, aliases in _ALIASES.items():
        header = _header_for(headers, aliases)
        if header:
            mapping[field_name] = header

    if "elan" in lowered_filename:
        villa = _header_for(headers, {"villanumber"}, last=True)
        if villa:
            mapping["unit_number"] = villa

    header_tokens = {_header_token(header.key): header.key for header in headers}
    if (
        "procedure" in header_tokens
        and "procedurepartytypenameen" not in header_tokens
        and "partytype" not in header_tokens
    ):
        mapping["party_role"] = header_tokens["procedure"]
        mapping.pop("procedure_type", None)

    family = "transaction_party" if mapping.get("party_role") else "contact_list"
    if mapping.get("transaction_date") and not mapping.get("party_role"):
        warnings.append("Dated owner records have no party role; ownership requires review.")
    if "contact_name" not in mapping:
        warnings.append("No owner/contact name column was detected.")
    if "unit_number" not in mapping:
        warnings.append("No direct unit-number column was detected.")
    if not {"phone_primary", "phone_alternate", "email_primary"} & mapping.keys():
        warnings.append("No contact channel was detected.")
    if family == "contact_list" and not mapping.get("building_name"):
        inferred = _inferred_project(filename, sheet_name)
        if inferred:
            warnings.append(f"Property name will be inferred as {inferred}.")
    return mapping, family, warnings


def _inferred_project(filename: str, sheet_name: str) -> str:
    source = f"{filename} {sheet_name}".lower()
    for name in ("Aura", "Harmony", "Alaya", "Elysian", "Elan"):
        if name.lower() in source:
            return name
    return ""


def _populated_xlsx_rows(workbook, sheet) -> list[int]:
    """Ignore formatting-only rows without stopping at gaps in real data.

    Scan XML once, then bound openpyxl's expensive cell/date decoding. Clear the
    sheetData parent too, so a million empty row elements do not accumulate.
    """
    rows = []
    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    parent = None
    with workbook._archive.open(sheet._worksheet_path) as stream:
        for action, element in ET.iterparse(stream, events=("start", "end")):
            if action == "start" and element.tag == ns + "sheetData":
                parent = element
            if action != "end" or element.tag != ns + "row":
                continue
            if any(node.text for node in element.iter() if node.tag in {ns + "v", ns + "t"}):
                rows.append(int(element.attrib["r"]))
            element.clear()
            if parent is not None:
                parent.clear()
    return rows


def _inspect_xlsx(data: bytes, filename: str) -> list[SheetProfile]:
    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    profiles: list[SheetProfile] = []
    try:
        for sheet in workbook.worksheets:
            initial_rows = list(
                sheet.iter_rows(
                    min_row=1,
                    max_row=min(15, max(1, sheet.max_row)),
                    max_col=min(100, max(1, sheet.max_column)),
                    values_only=True,
                )
            )
            header_row, headers = _find_header(initial_rows)
            mapping, family, warnings = _heuristic_mapping(
                headers,
                filename=filename,
                sheet_name=sheet.title,
            )
            populated_rows = _populated_xlsx_rows(workbook, sheet)
            profiles.append(
                SheetProfile(
                    name=sheet.title.strip() or "Sheet",
                    header_row=header_row,
                    headers=headers,
                    meaningful_rows=sum(number > header_row for number in populated_rows),
                    field_map=mapping,
                    sheet_family=family,
                    warnings=warnings,
                    last_data_row=max(populated_rows, default=header_row),
                )
            )
    finally:
        workbook.close()
    return profiles


def _csv_rows(data: bytes) -> list[list[str]]:
    text = data.decode("utf-8-sig")
    return list(csv.reader(io.StringIO(text)))


def _inspect_csv(data: bytes, filename: str) -> list[SheetProfile]:
    rows = _csv_rows(data)
    initial = [tuple(row[:100]) for row in rows[:15]]
    header_row, headers = _find_header(initial)
    mapping, family, warnings = _heuristic_mapping(
        headers,
        filename=filename,
        sheet_name=Path(filename).stem,
    )
    return [
        SheetProfile(
            name=Path(filename).stem,
            header_row=header_row,
            headers=headers,
            meaningful_rows=max(0, len(rows) - header_row),
            field_map=mapping,
            sheet_family=family,
            warnings=warnings,
        )
    ]


_DATE_TIME_SUFFIX = re.compile(
    r"[T\s]+\d{1,2}:\d{2}(?::\d{2}(?:\.\d+)?)?(?:\s*[ap]\.?m\.?)?$", re.IGNORECASE
)
_NUMERIC_DATE = re.compile(r"^(\d{1,2})[/-](\d{1,2})[/-](\d{4})$")
_DATE_FORMATS = (
    "%Y-%m-%d",
    "%Y/%m/%d",
    "%d %b %Y",
    "%d-%b-%Y",
    "%d %B %Y",
    "%d-%B-%Y",
    "%b %d %Y",
    "%B %d %Y",
)
_EMPTY_DATES = {"", "-", "n/a", "na", "none", "null", "unknown", "nan"}


def _date_text(value: Any) -> str:
    """Date text without a time part, with "Sept" as "Sep" and no commas."""
    text = _DATE_TIME_SUFFIX.sub("", _text(value)).strip()
    text = re.sub(r"\bsept\b", "Sep", text, flags=re.IGNORECASE)
    return " ".join(text.replace(",", " ").split())


def _numeric_date(text: str) -> tuple[int, int, int] | None:
    match = _NUMERIC_DATE.match(text)
    return (int(match[1]), int(match[2]), int(match[3])) if match else None


def _date_order(values: Iterable[Any]) -> str:
    """Day/month order of a date column: "dmy" (UAE, the default), "mdy" when a text date is
    clearly month first (09/25/2024), or "mixed" when the column has clear dates both ways."""
    day_first = month_first = False
    for value in values:
        if value is None or isinstance(value, date):
            continue
        parts = _numeric_date(_date_text(value))
        if not parts:
            continue
        first, second, _ = parts
        day_first = day_first or (first > 12 >= second)
        month_first = month_first or (second > 12 >= first)
        if day_first and month_first:
            return "mixed"
    return "mdy" if month_first else "dmy"


def _parse_date(value: Any, order: str = "dmy") -> date | None:
    """Read a sale date. NN/NN/YYYY follows its own clue when one number is over 12; otherwise
    the column's order, and a "mixed" column leaves it empty rather than guess."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _date_text(value)
    if not text:
        return None
    parts = _numeric_date(text)
    if parts:
        first, second, year = parts
        if first > 12 or (order == "dmy" and second <= 12):
            day, month = first, second
        elif second > 12 or order == "mdy" or first == second:
            day, month = second, first
        else:
            return None
        try:
            return date(year, month, day)
        except ValueError:
            return None
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _date_column_index(profile: SheetProfile) -> int | None:
    header_key = profile.field_map.get("transaction_date")
    header = next((item for item in profile.headers if item.key == header_key), None)
    return header.index if header_key and header else None


def _read_dates_by_column_order(profile: SheetProfile, values: Iterable[Any]) -> None:
    """Set the sheet's date order from its whole date column, and warn about dates read month
    first, left empty as unclear, or not readable at all."""
    header_key = profile.field_map.get("transaction_date")
    if not header_key:
        return
    values = list(values)
    profile.date_order = _date_order(values)
    label = next((h.label for h in profile.headers if h.key == header_key), header_key)
    unclear = unreadable = 0
    for value in values:
        if value is None or isinstance(value, date):
            continue
        if _text(value).lower() in _EMPTY_DATES or _parse_date(value, profile.date_order):
            continue
        parts = _numeric_date(_date_text(value))
        if parts and parts[0] <= 12 and parts[1] <= 12:
            unclear += 1
        else:
            unreadable += 1
    if profile.date_order == "mdy":
        profile.warnings.append(f"Dates in {label} were read month first (MM/DD/YYYY).")
    if unclear:
        profile.warnings.append(
            f"Dates in {label} mix day-first and month-first: {unclear:,} unclear "
            f"date{'s' if unclear != 1 else ''} left empty."
        )
    if unreadable:
        profile.warnings.append(
            f"{unreadable:,} date{'s' if unreadable != 1 else ''} in {label} could not be read."
        )


def _parse_amount(value: Any) -> int | None:
    if isinstance(value, int | float) and value > 0:
        return int(round(value))
    raw = re.sub(r"[^\d.]", "", _text(value))
    if not raw:
        return None
    try:
        amount = int(round(float(raw)))
    except ValueError:
        return None
    return amount if amount > 0 else None


def _normalized_role(value: Any) -> str:
    role = _text(value).lower()
    if "buyer" in role:
        return "Buyer"
    if "seller" in role:
        return "Seller"
    return _text(value).title()


def _field_value(
    row: tuple[Any, ...] | list[Any],
    profile: SheetProfile,
    canonical: str,
) -> Any:
    header_key = profile.field_map.get(canonical)
    if not header_key:
        return None
    header = next((item for item in profile.headers if item.key == header_key), None)
    if not header or header.index >= len(row):
        return None
    return row[header.index]


def _dropped_row(
    row: tuple[Any, ...] | list[Any],
    *,
    profile: SheetProfile,
    row_number: int,
    reason: str,
) -> dict[str, Any]:
    """The file's own values for a row the import does not keep, for the dropped-rows log."""

    def value(canonical: str) -> str:
        cell = _field_value(row, profile, canonical)
        if isinstance(cell, datetime):
            return cell.date().isoformat()
        if isinstance(cell, date):
            return cell.isoformat()
        return _text(cell)[:200]

    return {
        "sheet": profile.name,
        "row": row_number,
        "reason": reason,
        "name": value("contact_name"),
        "building": value("building_name") or value("project_name"),
        "unit": value("unit_number"),
        "date": value("transaction_date"),
        "price": value("transaction_price_aed"),
        "role": value("party_role"),
        "procedure": value("procedure_type"),
    }


def _missing_owner_name(name: str) -> bool:
    return name.strip().casefold() in {"", "-", "n/a", "na", "none", "null", "nan", "unknown"}


def _candidate_from_row(
    row: tuple[Any, ...] | list[Any],
    *,
    profile: SheetProfile,
    filename: str,
    row_number: int,
    dropped: list[dict[str, Any]] | None = None,
) -> ParsedContact | None:
    profile.property_inventory.add(
        {key: _text(_field_value(row, profile, key)) for key in profile.field_map}
    )
    name = _text(_field_value(row, profile, "contact_name"))
    unit = _normalize_unit(_field_value(row, profile, "unit_number"))
    if unit == "0":
        unit = ""
    unit_fallback_field = ""
    unit_fallback_value = ""
    if not unit:
        property_number = next(
            (
                _text(row[header.index])
                for header in profile.headers
                if _header_token(header.key) == "propertynumber"
                and header.index < len(row)
                and _text(row[header.index])
            ),
            "",
        )
        if property_number and _normalize_unit(property_number) not in {"0", "1"}:
            unit = _normalize_unit(property_number)
            unit_fallback_field = "propertynumber"
            unit_fallback_value = property_number
    building_no = _text(_field_value(row, profile, "building_number"))
    if _missing_owner_name(name):
        if dropped is not None:
            dropped.append(
                _dropped_row(row, profile=profile, row_number=row_number, reason="no_name")
            )
        return None
    role = _normalized_role(_field_value(row, profile, "party_role"))
    procedure = _text(_field_value(row, profile, "procedure_type"))
    inferred_project = _inferred_project(filename, profile.name)
    area = _text(_field_value(row, profile, "area_name"))
    master_project = _text(_field_value(row, profile, "master_project"))
    area_source_field = ""
    project = _text(_field_value(row, profile, "project_name")) or inferred_project
    building = _text(_field_value(row, profile, "building_name")) or project
    if not area:
        area = "Tilal Al Ghaf" if inferred_project else master_project
        if master_project and not inferred_project:
            area_source_field = "master_project"
    if not project:
        project = master_project

    phone = _normalize_phone(_field_value(row, profile, "phone_primary"))
    alternate = _normalize_phone(_field_value(row, profile, "phone_alternate"))
    channels = list(
        dict.fromkeys(
            value
            for value in (
                phone,
                alternate,
                *(
                    _normalize_phone(row[h.index])
                    for h in profile.headers
                    if h.index < len(row)
                    and _header_token(h.key)
                    in {
                        "mobile",
                        "mobile1",
                        "mobile2",
                        "phone",
                        "phone1",
                        "phone2",
                        "secondarymobile",
                        "alternatemobile",
                        "phonenumber",
                        "mobilenumber",
                    }
                ),
            )
            if value
        )
    )
    phone = channels[0] if channels else None
    alternate = channels[1] if len(channels) > 1 else None
    email = _normalize_email(_field_value(row, profile, "email_primary"))
    transaction_date = _parse_date(
        _field_value(row, profile, "transaction_date"), profile.date_order
    )
    transaction_price = _parse_amount(_field_value(row, profile, "transaction_price_aed"))
    fingerprint_source = "|".join(
        (
            profile.name,
            str(row_number),
            name.lower(),
            building.lower(),
            unit.lower(),
            transaction_date.isoformat() if transaction_date else "",
        )
    )
    candidate = ParsedContact(
        source_sheet=profile.name,
        source_row_number=row_number,
        contact_name=name[:300],
        phone_primary=phone,
        phone_alternate=alternate if alternate != phone else None,
        email_primary=email,
        area_name=area[:200],
        project_name=project[:200],
        building_name=building[:250],
        unit_number=unit[:100],
        property_type=_text(_field_value(row, profile, "property_type"))[:100],
        transaction_date=transaction_date,
        transaction_price_aed=transaction_price,
        party_role=role[:80],
        procedure_type=procedure[:160],
        row_fingerprint=hashlib.sha256(fingerprint_source.encode()).hexdigest(),
    )
    candidate.evidence.update(
        {
            "resolver_version": IMPORT_RESOLVER_VERSION,
            "source_filename": filename,
            "event_family": event_family(procedure),
            "source_fields": {
                canonical: _text(_field_value(row, profile, canonical))
                for canonical in profile.field_map
            },
            "contact_channels": channels,
            "source_identifiers": {
                key: _text(_field_value(row, profile, key))
                for key in (
                    "land_number",
                    "land_sub_number",
                    "building_number",
                    "property_number",
                    "plot_number",
                    "municipality_number",
                    "municipality_sub_number",
                )
            },
        }
    )
    if area_source_field:
        candidate.evidence["source_area_field"] = area_source_field
    source_unit = re.sub(r"\s+", "", _text(_field_value(row, profile, "unit_number"))).upper()
    if source_unit and source_unit.lower() != unit.lower():
        candidate.evidence["source_unit_number"] = source_unit[:100]
    if unit_fallback_field:
        compact_value = re.sub(r"\s+", "", unit_fallback_value).upper()
        candidate.evidence["source_unit_number"] = compact_value[:100]
        candidate.evidence["unit_number_fallback"] = {
            "source_field": unit_fallback_field,
            "source_value": unit_fallback_value[:100],
        }
    _recover_embedded_property_identity(candidate)
    address_recovery = candidate.evidence.get("address_recovery")
    building_no_identifier = _BUILDING_NO_NUMERIC_IDENTIFIER.fullmatch(building_no)
    if (
        not unit
        and isinstance(address_recovery, dict)
        and address_recovery.get("method") == "embedded_numeric_identifier"
        and building_no_identifier
    ):
        source_building = _text(_field_value(row, profile, "building_name"))
        candidate.building_name = _text(building_no_identifier.group(1))
        candidate.unit_number = _normalize_unit(building_no_identifier.group(2))
        candidate.evidence["source_unit_number"] = _normalize_unit(building_no)[:100]
        candidate.evidence["address_recovery"] = {
            "source_building_name": source_building,
            "source_building_no": building_no,
            "resolved_building_name": candidate.building_name,
            "resolved_unit_number": candidate.unit_number,
            "method": "building_no_structured_numeric_identifier",
        }
    # Keep the full code when the registry-style numeric suffix rule does not apply.
    if not candidate.unit_number and _BUILDING_NO_PROPERTY_IDENTIFIER.search(building_no):
        candidate.unit_number = _normalize_unit(building_no)[:100]
        candidate.evidence["source_unit_number"] = re.sub(r"\s+", "", building_no).upper()[:100]
    if (
        not candidate.unit_number
        and _normalize_unit(building_no) not in {"", "0", "1"}
        and any(character.isdigit() for character in building_no)
    ):
        candidate.unit_number = _normalize_unit(building_no)[:100]
        candidate.evidence["source_unit_number"] = re.sub(r"\s+", "", building_no).upper()[:100]
        candidate.evidence["unit_number_fallback"] = {
            "source_field": "buildingno",
            "source_value": building_no[:100],
        }
    return candidate


def _parse_xlsx(
    data: bytes,
    filename: str,
    profiles: list[SheetProfile],
    dropped: list[dict[str, Any]] | None = None,
) -> tuple[list[ParsedContact], int]:
    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    candidates: list[ParsedContact] = []
    total_rows = 0
    profile_by_name = {profile.name: profile for profile in profiles}
    try:
        for sheet in workbook.worksheets:
            profile = profile_by_name.get(sheet.title.strip() or "Sheet")
            if not profile:
                continue
            profile.property_inventory = SourcePropertyInventory()
            max_index = max((header.index for header in profile.headers), default=0)
            date_index = _date_column_index(profile)
            if date_index is not None:
                _read_dates_by_column_order(
                    profile,
                    (
                        row[0] if row else None
                        for row in sheet.iter_rows(
                            min_row=profile.header_row + 1,
                            max_row=profile.last_data_row,
                            min_col=date_index + 1,
                            max_col=date_index + 1,
                            values_only=True,
                        )
                    ),
                )
            for row_number, row in enumerate(
                sheet.iter_rows(
                    min_row=profile.header_row + 1,
                    max_row=profile.last_data_row,
                    max_col=max_index + 1,
                    values_only=True,
                ),
                start=profile.header_row + 1,
            ):
                if not any(value not in (None, "") for value in row):
                    continue
                total_rows += 1
                if total_rows > _MAX_ROWS:
                    raise ValueError(f"Import exceeds the {_MAX_ROWS:,}-row limit")
                candidate = _candidate_from_row(
                    row,
                    profile=profile,
                    filename=filename,
                    row_number=row_number,
                    dropped=dropped,
                )
                if candidate:
                    candidates.append(candidate)
    finally:
        workbook.close()
    return candidates, total_rows


def _parse_csv(
    data: bytes,
    filename: str,
    profiles: list[SheetProfile],
    dropped: list[dict[str, Any]] | None = None,
) -> tuple[list[ParsedContact], int]:
    rows = _csv_rows(data)
    profile = profiles[0]
    profile.property_inventory = SourcePropertyInventory()
    date_index = _date_column_index(profile)
    if date_index is not None:
        _read_dates_by_column_order(
            profile,
            (
                row[date_index] if date_index < len(row) else None
                for row in rows[profile.header_row :]
            ),
        )
    candidates: list[ParsedContact] = []
    total_rows = 0
    for row_number, row in enumerate(rows[profile.header_row :], start=profile.header_row + 1):
        if not any(value.strip() for value in row):
            continue
        total_rows += 1
        if total_rows > _MAX_ROWS:
            raise ValueError(f"Import exceeds the {_MAX_ROWS:,}-row limit")
        candidate = _candidate_from_row(
            row,
            profile=profile,
            filename=filename,
            row_number=row_number,
            dropped=dropped,
        )
        if candidate:
            candidates.append(candidate)
    return candidates, total_rows


def _summary(
    candidates: list[ParsedContact],
    total_rows: int,
    duplicate_rows_collapsed: int = 0,
    dropped_no_unit: int = 0,
) -> CrmContactImportSummary:
    return CrmContactImportSummary(
        total_rows=total_rows,
        candidate_rows=len(candidates),
        rows_with_sales_history=sum(c.latest_sale_date is not None for c in candidates),
        plot_only_rows=sum(
            bool(c.evidence.get("cadastral_identity")) and not c.unit_number for c in candidates
        ),
        distinct_properties=len(
            {candidate.property_key for candidate in candidates if candidate.property_key}
        ),
        duplicate_rows_collapsed=duplicate_rows_collapsed,
        dropped_no_unit=dropped_no_unit,
        valid_contacts=sum(candidate.has_valid_contact for candidate in candidates),
        verified_current_owners=sum(
            candidate.ownership_status == "verified_current_owner" for candidate in candidates
        ),
        probable_current_owners=sum(
            candidate.ownership_status == "probable_current_owner" for candidate in candidates
        ),
        unit_linked_unverified=sum(
            candidate.ownership_status == "unit_linked_unverified" for candidate in candidates
        ),
        former_owners=sum(candidate.ownership_status == "former_owner" for candidate in candidates),
        conflicting_claims=sum(
            candidate.ownership_status == "conflicting_claim" for candidate in candidates
        ),
        insufficient_property_data=sum(
            candidate.ownership_status == "insufficient_property_data" for candidate in candidates
        ),
        registry_not_observed=sum(
            candidate.ownership_status == "registry_not_observed" for candidate in candidates
        ),
        unmatched=sum(candidate.ownership_status == "unmatched" for candidate in candidates),
        rental_opportunities=sum(bool(candidate.lead_id) for candidate in candidates),
        contact_ready=sum(
            bool(candidate.lead_id)
            and candidate.has_valid_contact
            and candidate.ownership_status in {"verified_current_owner", "probable_current_owner"}
            for candidate in candidates
        ),
    )


def _mapping_response(profile: SheetProfile) -> dict[str, Any]:
    return CrmSheetMapping(
        sheet_name=profile.name,
        header_row=profile.header_row,
        meaningful_rows=profile.meaningful_rows,
        sheet_family=profile.sheet_family,
        mapping_source=profile.mapping_source,  # type: ignore[arg-type]
        confidence=profile.confidence,
        field_map=profile.field_map,
        warnings=profile.warnings,
    ).model_dump(mode="json")
